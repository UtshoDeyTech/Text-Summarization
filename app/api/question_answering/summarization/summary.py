from fastapi import APIRouter, HTTPException
from typing import List, Tuple, Optional
from pydantic import BaseModel, Field
import time
import tiktoken
from openai import AsyncOpenAI
import PyPDF2
import io
import json
from docx import Document
from app.service.log_client import logger
from config import OPENAI_API_KEY, MODEL, S3_BUCKET_NAME
from app.service import s3_storage
from pydantic import validator
from enum import Enum

router = APIRouter()
client = AsyncOpenAI(api_key=OPENAI_API_KEY)

MAX_TOKENS_PER_REQUEST = 3000
MAX_OUTPUT_TOKENS = 2000
BATCH_SIZE = 5
TOKEN_BUFFER = 100
MAX_PAGES = 10

class DocumentType(str, Enum):
    PDF = "pdf"
    DOC = "doc"
    DOCX = "docx"

class S3URLInput(BaseModel):
    url: str = Field(description="S3 URL of the document file to process")
    
    @validator('url')
    def validate_s3_url(cls, v):
        if not v.startswith(f"https://{S3_BUCKET_NAME}.s3."):
            raise ValueError("Invalid S3 URL format")
        
        lower_url = v.lower()
        if not any(lower_url.endswith(f".{ext}") for ext in [e.value for e in DocumentType]):
            raise ValueError("URL must point to a PDF, DOC, or DOCX file")
        return v

    def get_document_type(self) -> DocumentType:
        lower_url = self.url.lower()
        for doc_type in DocumentType:
            if lower_url.endswith(f".{doc_type.value}"):
                return doc_type
        raise ValueError("Unsupported document type")

class TokenUsage(BaseModel):
    prompt_tokens: int = Field(ge=0)
    completion_tokens: int = Field(ge=0)
    total_tokens: int = Field(ge=0)

class SummaryResponse(BaseModel):
    summary: str = Field(description="Generated HTML table summary")
    Carrier_name: str = Field(default="")
    Expiry_date: str = Field(default="")
    total_pages: int = Field(ge=0)
    pages_processed: int = Field(ge=0)
    batches_processed: int = Field(ge=0)
    summarization_levels: int = Field(ge=1, le=3)
    execution_time: float = Field(ge=0)
    token_usage: TokenUsage
    document_type: DocumentType
    status_code: str = Field(default="200")

class ErrorDetail(BaseModel):
    status_code: str
    error_messages: List[str]
    execution_time: Optional[float] = Field(None, ge=0)

class BatchSummaryResult(BaseModel):
    content: str
    success: bool
    token_usage: TokenUsage

def get_summary_prompt(level: int = 1, is_final: bool = False) -> str:
    if is_final:
        return """Extract and format insurance policy information in this exact JSON structure:
{
    "summary": "<table>...detailed HTML table with policy information...</table>",
    "Carrier_name": "insurance company/broker name - REQUIRED",
    "Expiry_date": "policy expiry date" 
}

TABLE REQUIREMENTS:
1. Document Information section must include:
   - Named Insured/Entity
   - Address details
   - Insurance Company/Broker name
   - Policy Number
   - Policy Period/Effective Dates
   - Policy Type
   - Form Numbers
   - Premium Details

2. Content Details section must include:
   - Coverage Types and Limits
   - Deductibles
   - Policy Terms
   - Conditions
   - Exclusions
   - Claims Procedures
   - Cancellation Terms
   - Notice Requirements
   - Special Provisions
   - Premium Payment Terms
   - Endorsements

Table Format:
<table>
<tr><th>Field</th><th>Details</th></tr>
<tr><th colspan="2">Document Information</th></tr>
<tr><td>[Insurance Details]</td><td>[Value]</td></tr>
<tr><th colspan="2">Content Details</th></tr>
<tr><td>[Policy Details]</td><td>[Value]</td></tr>
</table>

CRITICAL RULES:
1. ALWAYS extract and include insurance company/broker name in Carrier_name field
2. Look for carrier name in:
   - Insurance Company field
   - Broker/Agent field
   - Declarations page
   - Header/Footer
   - Company logos/letterhead
3. If multiple companies listed, use the primary carrier
4. For expiry date, check:
   - Policy Period End Date
   - Expiration Date
   - Policy Term End
   - Renewal Date
   
Format Rules:
- Include only verified information
- Use "Not specified" for missing data
- No empty cells/rows
- No duplicate information
- No line breaks in cells
- Clean single-line HTML"""
    else:
        return f"""Extract Level {level} insurance policy information:
1. Core insurance details:
   - Carrier/Broker name (CRITICAL)
   - Policy dates and numbers
   - Insured details
   - Premium information

2. Policy specifics:
   - Coverage details
   - Terms and conditions
   - Claims procedures
   - Important provisions"""

def count_tokens(text: str) -> int:
    enc = tiktoken.encoding_for_model(MODEL)
    return len(enc.encode(text))

def truncate_to_token_limit(text: str, max_tokens: int) -> str:
    enc = tiktoken.encoding_for_model(MODEL)
    tokens = enc.encode(text)
    if len(tokens) <= max_tokens:
        return text
    return enc.decode(tokens[:max_tokens])

def extract_broker_from_summary(summary: str) -> str:
    """Extract broker name from summary table if present"""
    if 'Broker' in summary:
        try:
            broker_start = summary.index('Broker</td><td>') + len('Broker</td><td>')
            broker_end = summary.index('</td>', broker_start)
            return summary[broker_start:broker_end]
        except ValueError:
            pass
    return "Not specified"

def ensure_complete_table_structure(table_html: str) -> str:
    """Ensure table has all required sections"""
    if not table_html.startswith('<table>'):
        table_html = f'<table>{table_html}'
    if not table_html.endswith('</table>'):
        table_html = f'{table_html}</table>'
        
    sections = ['Document Information', 'Content Details']
    for section in sections:
        if section not in table_html:
            insert_point = table_html.rfind('</table>')
            section_html = f'<tr><th colspan="2">{section}</th></tr><tr><td>Information</td><td>Not specified</td></tr>'
            table_html = table_html[:insert_point] + section_html + table_html[insert_point:]
            
    return table_html

async def create_safe_batch_summary(chunks: List[str], level: int) -> BatchSummaryResult:
    max_tokens_per_chunk = (MAX_TOKENS_PER_REQUEST - TOKEN_BUFFER) // len(chunks)
    truncated_chunks = [truncate_to_token_limit(chunk, max_tokens_per_chunk) for chunk in chunks]
    combined_text = "\n\n".join([f"Chunk {i+1}:\n{chunk}" for i, chunk in enumerate(truncated_chunks)])
    
    try:
        response = await client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": get_summary_prompt(level)},
                {"role": "user", "content": f"Extract and organize key information:\n\n{combined_text}"}
            ],
            temperature=0.7,
            max_tokens=MAX_OUTPUT_TOKENS
        )

        return BatchSummaryResult(
            content=response.choices[0].message.content,
            success=True,
            token_usage=TokenUsage(
                prompt_tokens=response.usage.prompt_tokens,
                completion_tokens=response.usage.completion_tokens,
                total_tokens=response.usage.total_tokens
            )
        )
    except Exception as e:
        logger.error(f"Error in batch summary: {str(e)}")
        return BatchSummaryResult(
            content="",
            success=False,
            token_usage=TokenUsage(prompt_tokens=0, completion_tokens=0, total_tokens=0)
        )

async def process_chunks_hierarchically(chunks: List[str]) -> Tuple[str, int, int, TokenUsage, str, str]:
    current_chunks = chunks
    level = 1
    batch_count = 0
    total_token_usage = TokenUsage(prompt_tokens=0, completion_tokens=0, total_tokens=0)
    
    while len(current_chunks) > BATCH_SIZE:
        logger.info(f"Processing level {level} with {len(current_chunks)} chunks")
        next_level_chunks = []
        
        for i in range(0, len(current_chunks), BATCH_SIZE):
            batch = current_chunks[i:i + BATCH_SIZE]
            batch_count += 1
            
            result = await create_safe_batch_summary(batch, level)
            if not result.success:
                raise HTTPException(
                    status_code=500,
                    detail=ErrorDetail(
                        status_code="500",
                        error_messages=["Failed to process batch"]
                    ).dict()
                )

            total_token_usage.prompt_tokens += result.token_usage.prompt_tokens
            total_token_usage.completion_tokens += result.token_usage.completion_tokens
            total_token_usage.total_tokens += result.token_usage.total_tokens
            next_level_chunks.append(result.content)
        
        current_chunks = next_level_chunks
        level += 1
        
        if level > 3:
            logger.warning("Reached maximum summarization levels")
            break
    
    try:
        final_response = await client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": get_summary_prompt(is_final=True)},
                {"role": "user", "content": "Extract and summarize the key information from these documents into the required JSON format:\n\n" + "\n\n".join(current_chunks)}
            ],
            temperature=0.7,
            max_tokens=2000,
            response_format={"type": "json_object"}
        )

        total_token_usage.prompt_tokens += final_response.usage.prompt_tokens
        total_token_usage.completion_tokens += final_response.usage.completion_tokens
        total_token_usage.total_tokens += final_response.usage.total_tokens

        try:
            response_data = json.loads(final_response.choices[0].message.content)
            
            final_summary = response_data.get('summary', '')
            if not final_summary.strip():
                raise ValueError("Empty summary in response")
                
            carrier_name = response_data.get('Carrier_name', '')
            if carrier_name == "Not specified" or not carrier_name:
                if "AMWINS" in final_summary:
                    carrier_name = "AMWINS INS BROKERAGE LLC"
                elif "Broker" in final_summary:
                    carrier_name = extract_broker_from_summary(final_summary)
                    
            expiry_date = response_data.get('Expiry_date', '')
            
            if not all(section in final_summary for section in ['Document Information', 'Content Details']):
                logger.warning("Missing required sections in summary")
                final_summary = ensure_complete_table_structure(final_summary)
            
        except (ValueError, json.JSONDecodeError) as json_error:
            logger.error(f"Error parsing JSON response: {str(json_error)}")
            raise HTTPException(
                status_code=500,
                detail=ErrorDetail(
                    status_code="500",
                    error_messages=["Failed to parse summary response"]
                ).dict()
            )

    except Exception as e:
        logger.error(f"Error in final formatting: {str(e)}")
        raise HTTPException(
            status_code=500,
            detail=ErrorDetail(
                status_code="500",
                error_messages=["Failed to create summary"]
            ).dict()
        )
    
    return final_summary, batch_count, level, total_token_usage, carrier_name, expiry_date

async def get_document_from_s3(url: str) -> bytes:
    try:
        from urllib.parse import unquote
        base_url = f"https://{S3_BUCKET_NAME}.s3.us-east-1.amazonaws.com/"
        if base_url not in url:
            raise ValueError("Invalid S3 URL format")
            
        object_key = unquote(url.replace(base_url, ""))
        logger.info(f"Extracted object key: {object_key}")
        
        if not s3_storage.verify_bucket(S3_BUCKET_NAME):
            raise HTTPException(status_code=404, detail=ErrorDetail(
                status_code="404",
                error_messages=["S3 bucket not accessible"]
            ).dict())
        
        response = s3_storage.s3_client.get_object(Bucket=S3_BUCKET_NAME, Key=object_key)
        return response['Body'].read()
        
    except ValueError as e:
        logger.error(f"URL format error: {str(e)}")
        raise HTTPException(status_code=400, detail=ErrorDetail(
            status_code="400",
            error_messages=[str(e)]
        ).dict())
    except s3_storage.ClientError as e:
        logger.error(f"S3 client error: {str(e)}")
        raise HTTPException(status_code=404, detail=ErrorDetail(
            status_code="404",
            error_messages=[f"S3 error: {str(e)}"]
        ).dict())
    except Exception as e:
        logger.error(f"Error fetching from S3: {str(e)}")
        raise HTTPException(status_code=500, detail=ErrorDetail(
            status_code="500",
            error_messages=[f"Failed to fetch from S3: {str(e)}"]
        ).dict())

def extract_text_from_pdf(pdf_file: bytes) -> Tuple[List[str], int]:
    try:
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(pdf_file))
        total_pages = len(pdf_reader.pages)
        pages_to_process = min(total_pages, MAX_PAGES)
        
        chunks = []
        for page_num in range(pages_to_process):
            text = pdf_reader.pages[page_num].extract_text().strip()
            if text:
                chunks.append(text)
        
        return chunks, total_pages
    except Exception as e:
        logger.error(f"Error processing PDF: {str(e)}")
        raise HTTPException(status_code=400, detail=ErrorDetail(
            status_code="400",
            error_messages=["Failed to process PDF file"]
        ).dict())

def extract_text_from_docx(doc_file: bytes) -> Tuple[List[str], int]:
    try:
        doc = Document(io.BytesIO(doc_file))
        total_pages = len(doc.paragraphs) // 40
        
        chunks = []
        current_chunk = []
        current_length = 0
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
                
            current_chunk.append(text)
            current_length += len(text)
            
            if current_length > 2000:
                chunks.append("\n".join(current_chunk))
                current_chunk = []
                current_length = 0
        
        if current_chunk:
            chunks.append("\n".join(current_chunk))
        
        return chunks, max(1, total_pages)
    except Exception as e:
        logger.error(f"Error processing DOCX: {str(e)}")
        raise HTTPException(status_code=400, detail=ErrorDetail(
            status_code="400",
            error_messages=["Failed to process DOCX file"]
        ).dict())

@router.post("/summarize/document", response_model=SummaryResponse)
async def summarize_document(input_data: S3URLInput) -> SummaryResponse:
    start_time = time.time()
    
    try:
        document_type = input_data.get_document_type()
        document_content = await get_document_from_s3(input_data.url)
        
        if document_type == DocumentType.PDF:
            chunks, total_pages = extract_text_from_pdf(document_content)
        elif document_type in [DocumentType.DOC, DocumentType.DOCX]:
            chunks, total_pages = extract_text_from_docx(document_content)
        else:
            raise HTTPException(status_code=400, detail=ErrorDetail(
                status_code="400",
                error_messages=["Unsupported document type"]
            ).dict())
        
        if not chunks:
            raise HTTPException(status_code=400, detail=ErrorDetail(
                status_code="400",
                error_messages=["No readable text found in document"]
            ).dict())
        
        logger.info(f"Processing {len(chunks)} chunks from {document_type.value} document")
        final_summary, batches_processed, levels, token_usage, carrier_name, expiry_date = await process_chunks_hierarchically(chunks)
        
        return SummaryResponse(
            summary=final_summary,
            Carrier_name=carrier_name,
            Expiry_date=expiry_date,
            total_pages=total_pages,
            pages_processed=min(total_pages, MAX_PAGES),
            batches_processed=batches_processed,
            summarization_levels=levels,
            execution_time=round(time.time() - start_time, 2),
            token_usage=token_usage,
            document_type=document_type,
            status_code="200"
        )
        
    except HTTPException:
        raise
    except Exception as e:
        error_msg = f"Error generating summary: {str(e)}"
        logger.error(error_msg)
        raise HTTPException(status_code=500, detail=ErrorDetail(
            status_code="500",
            error_messages=[str(e)],
            execution_time=round(time.time() - start_time, 2)
        ).dict())