from fastapi import APIRouter, UploadFile, File, HTTPException, Request, Form
from fastapi.responses import JSONResponse
from typing import Optional
from datetime import datetime
from PyPDF2 import PdfReader
from docx import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from io import BytesIO
from tqdm import tqdm
from app.service.openai_client import get_embeddings
from app.service.pinecone_client import (
    upsert_vectors,
    find_document_namespace,
    delete_vectors
)
from app.service.log_client import logger
from app.service.s3_storage import s3_client, verify_bucket
from botocore.exceptions import ClientError
from config import S3_BUCKET_NAME
import time

router = APIRouter()
text_splitter = RecursiveCharacterTextSplitter(chunk_size=1500, chunk_overlap=100)

SUPPORTED_EXTENSIONS = {
    'pdf': 'application/pdf',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
}

def truncate_text_for_metadata(text: str, max_bytes: int = 35000) -> str:
    """Truncate text to stay within metadata size limits"""
    if not text:
        return ""
    encoded = text.encode('utf-8')
    if len(encoded) <= max_bytes:
        return text
    return encoded[:max_bytes].decode('utf-8', 'ignore')

def extract_text_from_pdf(file_obj):
    pdf = PdfReader(file_obj)
    text = ""
    
    with tqdm(total=len(pdf.pages), desc="Processing PDF pages") as pbar:
        for i, page in enumerate(pdf.pages):
            text += page.extract_text()
            pbar.update(1)
            logger.info(f"Extracted text from PDF page {i+1}/{len(pdf.pages)}")
    
    return text

def extract_text_from_docx(file_obj):
    doc = Document(file_obj)
    text = ""
    
    with tqdm(total=len(doc.paragraphs), desc="Processing DOCX paragraphs") as pbar:
        for para in doc.paragraphs:
            text += para.text + "\n"
            pbar.update(1)
    
    return text

def process_document(file_obj, file_extension):
    try:
        logger.info(f"Starting text extraction | file_type={file_extension}")
        
        if file_extension == 'pdf':
            text = extract_text_from_pdf(file_obj)
            chunks = text_splitter.split_text(text)
        elif file_extension == 'docx':
            text = extract_text_from_docx(file_obj)
            chunks = text_splitter.split_text(text)
        else:
            raise ValueError(f"Unsupported file type: {file_extension}")

        if not chunks:
            logger.warning("No text extracted from document")
            return []

        final_chunks = []
        for chunk in chunks:
            truncated_chunk = truncate_text_for_metadata(chunk)
            if truncated_chunk:
                final_chunks.append(truncated_chunk)

        logger.info(f"Document chunking completed | total_chunks={len(final_chunks)}")
        return final_chunks
        
    except Exception as e:
        logger.error(f"Document processing failed | error_type={type(e).__name__}, error={str(e)}")
        raise

@router.post("/document_url_upload/{document_id}")
async def document_url_upload(
    request: Request,
    document_id: str,
    document_category: Optional[str] = "Other",
    file: Optional[UploadFile] = File(None),
    url: Optional[str] = Form(None)
):
    start_time = datetime.utcnow()
    
    if file and url:
        raise HTTPException(
            status_code=400,
            detail={
                "status_code": "400",
                "error_messages": ["Only one of file or URL parameters is allowed"]
            }
        )
    
    # Determine input type (file or URL)
    using_file = file is not None
    document_url = url if not using_file else None
    
    if not using_file and not document_url:
        raise HTTPException(
            status_code=400,
            detail={
                "status_code": "400",
                "error_messages": ["Either file or URL parameter is required"]
            }
        )
    
    try:
        file_obj = None
        file_extension = None
        filename = None
        
        if using_file:
            # Process uploaded file
            filename = file.filename
            file_extension = filename.split('.')[-1].lower()
            
            if file_extension not in SUPPORTED_EXTENSIONS:
                raise HTTPException(
                    status_code=400,
                    detail={
                        "status_code": "400",
                        "error_messages": [f"Supported file types are: {', '.join(SUPPORTED_EXTENSIONS.keys())}"]
                    }
                )
            
            file_content = await file.read()
            file_obj = BytesIO(file_content)
            logger.info(f"Processing uploaded file | filename={filename}, size={len(file_content)/1024/1024:.2f} MB")
            
        else:
            # Process file from URL
            # Extract the file extension from the URL
            url_parts = document_url.split('/')
            filename = url_parts[-1]
            file_extension = filename.split('.')[-1].lower()
            
            if file_extension not in SUPPORTED_EXTENSIONS:
                raise HTTPException(
                    status_code=400,
                    detail={
                        "status_code": "400",
                        "error_messages": [f"Supported file types are: {', '.join(SUPPORTED_EXTENSIONS.keys())}"]
                    }
                )
                
            # Check if S3 bucket is valid
            if not verify_bucket(S3_BUCKET_NAME):
                raise HTTPException(
                    status_code=500,
                    detail={
                        "status_code": "500",
                        "error_messages": ["Could not access S3 bucket"]
                    }
                )
            
            # Extract the object key from the URL
            if '.amazonaws.com/' in document_url:
                object_key = document_url.split('.amazonaws.com/')[1]
            else:
                # If URL format is different, try to extract key differently
                object_key = '/'.join(url_parts[-3:])  # Assuming the last 3 segments form the key
            
            logger.info(f"Downloading file from S3 | bucket={S3_BUCKET_NAME}, key={object_key}")
            
            # Download the file from S3 with progress tracking
            file_obj = BytesIO()
            try:
                # Get file size first to log progress
                response = s3_client.head_object(Bucket=S3_BUCKET_NAME, Key=object_key)
                total_size = response.get('ContentLength', 0)
                logger.info(f"Starting download of {total_size/1024/1024:.2f} MB file")
                
                # Custom callback to track progress
                downloaded_bytes = 0
                last_log_time = time.time()
                
                def log_progress(bytes_transferred):
                    nonlocal downloaded_bytes, last_log_time
                    downloaded_bytes += bytes_transferred
                    current_time = time.time()
                    if current_time - last_log_time > 5:  # Log every 5 seconds
                        logger.info(f"Downloaded {downloaded_bytes/1024/1024:.2f} MB of {total_size/1024/1024:.2f} MB ({downloaded_bytes/total_size*100:.1f}%)")
                        last_log_time = current_time
                
                s3_client.download_fileobj(
                    S3_BUCKET_NAME, 
                    object_key, 
                    file_obj, 
                    Callback=log_progress
                )
                file_obj.seek(0)  # Reset file pointer to beginning
                logger.info(f"Download complete | file_size={total_size/1024/1024:.2f} MB")
                
            except ClientError as e:
                error_code = e.response.get('Error', {}).get('Code', 'Unknown')
                logger.error(f"S3 download failed | bucket={S3_BUCKET_NAME}, key={object_key}, error_code={error_code}, error={str(e)}")
                raise HTTPException(
                    status_code=404,
                    detail={
                        "status_code": "404",
                        "error_messages": ["File not found in S3 bucket"]
                    }
                )
        
        # Check if the document already exists in Pinecone
        namespace = find_document_namespace(document_id)
        if namespace:
            delete_vectors(document_id)
            logger.info(f"Deleted existing vectors from Pinecone | document_id={document_id}")
        
        # Process the document
        processing_start = time.time()
        logger.info("Starting document processing")
        chunks = process_document(file_obj, file_extension)
        processing_duration = time.time() - processing_start
        logger.info(f"Document processing completed | duration_seconds={processing_duration:.2f}, chunks={len(chunks)}")
        
        if len(chunks) == 0:
            raise HTTPException(
                status_code=400,
                detail={
                    "status_code": "400",
                    "error_messages": ["No text could be extracted from the document"]
                }
            )
        
        # Create embeddings
        embedding_start = time.time()
        logger.info(f"Starting embedding generation for {len(chunks)} chunks")
        with tqdm(total=len(chunks), desc="Creating embeddings") as pbar:
            embeddings = await get_embeddings(chunks)
            pbar.update(len(chunks))
        embedding_duration = time.time() - embedding_start
        logger.info(f"Embedding generation completed | duration_seconds={embedding_duration:.2f}, embeddings={len(embeddings)}")
        
        # Prepare vectors for Pinecone
        base_metadata = {
            "document_id": document_id,
            "filename": filename,
            "file_type": file_extension,
            "document_category": document_category, 
            "upload_date": datetime.utcnow().isoformat()
        }
        
        # Add source_url to metadata if using URL
        if not using_file and document_url:
            base_metadata["source_url"] = document_url
            
        base_size = len(str(base_metadata).encode('utf-8'))
        max_text_size = 35000 - base_size
        
        ids = []
        metadatas = []
        valid_chunks = []
        valid_embeddings = []
        
        for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
            truncated_text = truncate_text_for_metadata(chunk, max_text_size)
            if truncated_text:
                metadata = base_metadata.copy()
                metadata["text"] = truncated_text
                ids.append(f"{document_id}_{i}")
                metadatas.append(metadata)
                valid_chunks.append(chunk)
                valid_embeddings.append(embedding)
        
        # Upload to Pinecone
        pinecone_start = time.time()
        logger.info(f"Starting Pinecone upload | vectors={len(valid_embeddings)}")
        with tqdm(total=1, desc="Uploading to Pinecone") as pbar:
            upsert_vectors(valid_embeddings, metadatas, ids, document_id)
            pbar.update(1)
        pinecone_duration = time.time() - pinecone_start
        logger.info(f"Pinecone upload completed | duration_seconds={pinecone_duration:.2f}")
        
        # Log detailed performance metrics
        if not using_file:
            logger.info(f"Performance metrics | download_time={processing_start - start_time.timestamp():.2f}s, " +
                       f"processing_time={processing_duration:.2f}s, " +
                       f"embedding_time={embedding_duration:.2f}s, " + 
                       f"pinecone_time={pinecone_duration:.2f}s")
            if 'total_size' in locals():
                logger.info(f"File size: {total_size/1024/1024:.2f}MB")
        
        total_duration = (datetime.utcnow() - start_time).total_seconds()
        logger.info(f"Document processing workflow completed | document_id={document_id}, total_duration_seconds={total_duration}")
        
        # Return standardized response format
        return JSONResponse(
            content={
                "status": "success",
                "document_id": document_id,
                "chunks_stored": len(valid_chunks),
                "file_type": file_extension,
                "document_category": document_category,
                "processing_time_seconds": total_duration,
                "status_code": "200"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        source_info = f"url={document_url}" if not using_file else f"file={filename}"
        error_msg = f"Document upload failed | document_id={document_id}, {source_info}, error={str(e)}"
        logger.error(error_msg)
        raise HTTPException(
            status_code=500,
            detail={
                "status_code": "500",
                "error_messages": [str(e)]
            }
        )