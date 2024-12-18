import uuid
import logging
from datetime import datetime
from fastapi import APIRouter, UploadFile, File, Form, HTTPException, Request, BackgroundTasks
from fastapi.responses import JSONResponse
from PyPDF2 import PdfReader
from docx import Document
import pandas as pd
from langchain.text_splitter import RecursiveCharacterTextSplitter
from io import BytesIO
from botocore.exceptions import ClientError
from tqdm import tqdm
from openpyxl import load_workbook
from app.service.s3_storage import (
    upload_file, 
    delete_object, 
    s3_client, 
    S3_BUCKET_NAME
)
from app.service.openai_client import get_embeddings
from app.service.pinecone_client import (
    upsert_vectors,
    list_all_vectors,
    delete_vectors,
    find_document_namespace
)
from app.service.log_client import logger
import requests
from config import AI_VALUE_ASP

router = APIRouter()
text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)  # Reduced chunk size

SUPPORTED_EXTENSIONS = {
    'pdf': 'application/pdf',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'csv': 'text/csv',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'xls': 'application/vnd.ms-excel'
}

def truncate_text_for_metadata(text: str, max_bytes: int = 35000) -> str:
    """Truncate text to stay within metadata size limits"""
    if not text:
        return ""
    encoded = text.encode('utf-8')
    if len(encoded) <= max_bytes:
        return text
    return encoded[:max_bytes].decode('utf-8', 'ignore')

def create_doc_info(payload, headers):
    url = AI_VALUE_ASP+"/api/docinfo/createdocinfo"
    try:
        response = requests.post(url, json=payload, headers=headers)
        response_data = response.json()
        
        if response_data.get('status_code') != "200":
            error_message = response_data.get('error_messages', ['Something went wrong'])[0]
            raise Exception(error_message)
            
        return response_data
    except requests.exceptions.RequestException as e:
        logger.error(f"Error creating doc info: {e}")
        return None

def extract_text_from_pdf(file_obj):
    pdf = PdfReader(file_obj)
    text = ""
    
    with tqdm(total=len(pdf.pages), desc="Processing PDF pages") as pbar:
        for i, page in enumerate(pdf.pages):
            text += page.extract_text()
            pbar.update(1)
            logger.info(f"Extracted text from PDF page {i+1}/{len(pdf.pages)}")
    
    return text

def extract_text_from_docx(file_obj):
    doc = Document(file_obj)
    text = ""
    
    with tqdm(total=len(doc.paragraphs), desc="Processing DOCX paragraphs") as pbar:
        for para in doc.paragraphs:
            text += para.text + "\n"
            pbar.update(1)
    
    return text

def extract_text_from_csv(file_obj, chunk_size=1000):
    """Extract text from CSV file with chunked processing."""
    try:
        text_chunks = []
        total_rows = sum(1 for _ in pd.read_csv(file_obj, chunksize=chunk_size))
        file_obj.seek(0)
        
        with tqdm(total=total_rows, desc="Processing CSV rows") as pbar:
            for chunk in pd.read_csv(file_obj, chunksize=chunk_size):
                chunk = chunk.astype(str).apply(lambda x: x.str[:200])  # Limit cell content
                chunk_text = chunk.fillna('').to_string(index=False)
                text_chunks.append(chunk_text)
                pbar.update(len(chunk))
                
        return '\n'.join(text_chunks)
    except Exception as e:
        logger.error(f"Error extracting text from CSV: {e}")
        raise

def extract_text_from_excel(file_obj):
    """Excel processing with text size limits for Pinecone"""
    try:
        text_chunks = []
        
        xl = pd.ExcelFile(file_obj)
        sheet_names = xl.sheet_names
        logger.info(f"Found {len(sheet_names)} sheets in Excel file")
        
        for sheet_name in sheet_names:
            logger.info(f"Processing sheet: {sheet_name}")
            text_chunks.append(f"\n=== Sheet: {sheet_name} ===\n")
            
            # Read the entire sheet at once for better performance
            df = pd.read_excel(xl, sheet_name=sheet_name)
            total_rows = len(df)
            logger.info(f"Sheet {sheet_name} has {total_rows} rows")
            
            # Process in smaller chunks
            chunk_size = 50  # Smaller chunks for better metadata handling
            for start_idx in range(0, total_rows, chunk_size):
                end_idx = min(start_idx + chunk_size, total_rows)
                
                # Process chunk
                chunk_df = df.iloc[start_idx:end_idx].copy()
                # Convert to string and limit cell content
                chunk_df = chunk_df.astype(str).apply(lambda x: x.str[:200])
                chunk_text = chunk_df.fillna('').to_string(index=False)
                
                # Ensure chunk size is within limits
                chunk_text = truncate_text_for_metadata(chunk_text)
                if chunk_text:
                    text_chunks.append(chunk_text)
                
                logger.info(f"Processed rows {start_idx} to {end_idx} in sheet {sheet_name}")
            
            logger.info(f"Completed processing sheet {sheet_name} | total_rows={total_rows}")
            
        return text_chunks
    except Exception as e:
        logger.error(f"Error extracting text from Excel: {str(e)}")
        raise

def process_document(file_obj, file_extension):
    try:
        logger.info(f"Starting text extraction | file_type={file_extension}")
        
        if file_extension == 'pdf':
            text = extract_text_from_pdf(file_obj)
            chunks = text_splitter.split_text(text)
        elif file_extension == 'docx':
            text = extract_text_from_docx(file_obj)
            chunks = text_splitter.split_text(text)
        elif file_extension == 'csv':
            text = extract_text_from_csv(file_obj)
            chunks = text_splitter.split_text(text)
        elif file_extension in ['xlsx', 'xls']:
            chunks = extract_text_from_excel(file_obj)
        else:
            raise ValueError(f"Unsupported file type: {file_extension}")

        if not chunks:
            logger.warning("No text extracted from document")
            return []

        # Ensure chunks are within metadata limits
        final_chunks = []
        for chunk in chunks:
            truncated_chunk = truncate_text_for_metadata(chunk)
            if truncated_chunk:
                final_chunks.append(truncated_chunk)

        logger.info(f"Document chunking completed | total_chunks={len(final_chunks)}")
        return final_chunks
        
    except Exception as e:
        logger.error(f"Document processing failed | error_type={type(e).__name__}, error={str(e)}")
        raise

@router.post("/{user_id}/upload_document")
async def upload_document(
    request: Request,
    user_id: str,
    BEARER_TOKEN: str = Form(...),
    file: UploadFile = File(...)
):
    start_time = datetime.utcnow()
    try:
        file_extension = file.filename.split('.')[-1].lower()
        document_id = file.filename
        object_name = f"{user_id}/{document_id}"

        if file_extension not in SUPPORTED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail={
                    "status_code": "400",
                    "error_messages": [f"Supported file types are: {', '.join(SUPPORTED_EXTENSIONS.keys())}"]
                }
            )

        # Check if file exists in S3
        try:
            s3_client.head_object(Bucket=S3_BUCKET_NAME, Key=object_name)
            file_exists_s3 = True
        except ClientError:
            file_exists_s3 = False

        # Check if file exists in Pinecone
        namespace = find_document_namespace(user_id, document_id)
        file_exists_pinecone = namespace is not None

        # Delete existing file if it exists
        if file_exists_s3:
            delete_object(S3_BUCKET_NAME, object_name)
            logger.info(f"Deleted existing file from S3 | object={object_name}")
        if file_exists_pinecone:
            delete_vectors(user_id, document_id)
            logger.info(f"Deleted existing vectors from Pinecone | document_id={document_id}")

        # Upload to S3
        file_content = await file.read()
        s3_url = upload_file(BytesIO(file_content), S3_BUCKET_NAME, object_name)
        if not s3_url:
            raise HTTPException(
                status_code=500,
                detail={
                    "status_code": "500",
                    "error_messages": ["Failed to upload document to S3"]
                }
            )

        # Process document and create chunks
        chunks = process_document(BytesIO(file_content), file_extension)
        if len(chunks) == 0:
            raise HTTPException(
                status_code=400,
                detail={
                    "status_code": "400",
                    "error_messages": ["No text could be extracted from the document"]
                }
            )
        
        # Create embeddings
        with tqdm(total=len(chunks), desc="Creating embeddings") as pbar:
            embeddings = get_embeddings(chunks)
            pbar.update(len(chunks))

        # Create metadata with size checking
        ids = []
        metadatas = []
        valid_chunks = []
        valid_embeddings = []

        base_metadata = {
            "document_id": document_id,
            "user_id": user_id,
            "filename": file.filename,
            "file_type": file_extension,
            "docs_category": "",
            "org_id": "",
            "upload_date": datetime.utcnow().isoformat()
        }
        base_size = len(str(base_metadata).encode('utf-8'))
        max_text_size = 35000 - base_size

        for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
            truncated_text = truncate_text_for_metadata(chunk, max_text_size)
            if truncated_text:
                metadata = base_metadata.copy()
                metadata["text"] = truncated_text
                ids.append(f"{document_id}_{i}")
                metadatas.append(metadata)
                valid_chunks.append(chunk)
                valid_embeddings.append(embedding)

        # Upload to Pinecone
        with tqdm(total=1, desc="Uploading to Pinecone") as pbar:
            upsert_vectors(user_id, valid_embeddings, metadatas, ids, file.filename)
            pbar.update(1)

        pinecone_duration = (datetime.utcnow() - start_time).total_seconds()
        logger.info(f"Pinecone upload completed | document_id={document_id}, chunks={len(valid_chunks)}, duration_seconds={pinecone_duration}")

        # Create document info
        payload = {
            "userid": user_id,
            "name": document_id,
            "s3url": s3_url,
            "s3path": f"{S3_BUCKET_NAME}/{user_id}",
            "vectorid": f"document-vectors-{user_id}"
        }
        headers = {
            'accept': 'application/json',
            'Authorization': f'Bearer {BEARER_TOKEN}',
            'Content-Type': 'application/json'
        }

        response = create_doc_info(payload, headers)
        
        total_duration = (datetime.utcnow() - start_time).total_seconds()
        logger.info(f"Document processing completed | user_id={user_id}, document_id={document_id}, total_duration_seconds={total_duration}")

        return JSONResponse(
            content={
                "document_id": document_id,
                "user_id": user_id,
                "chunks_stored": len(valid_chunks),
                "s3_url": s3_url,
                "file_type": file_extension,
                "processing_time_seconds": total_duration,
                "response": response,
                "was_overwrite": file_exists_s3 or file_exists_pinecone,
                "status_code": "200"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        error_msg = f"Document upload failed | user_id={user_id}, filename={file.filename}, error={str(e)}"
        logger.error(error_msg)
        raise HTTPException(
            status_code=500,
            detail={
                "status_code": "500",
                "error_messages": [str(e)]
            }
        )